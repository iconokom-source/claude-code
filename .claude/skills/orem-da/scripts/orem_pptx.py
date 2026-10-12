#!/usr/bin/env python3
"""Génère une présentation PowerPoint (.pptx) aux couleurs d'Orem.

Les coordonnées sont exprimées en pixels d'une slide Figma 1920×1080
(1 px = 6350 EMU, 1 px de police = 0,5 pt), ce qui permet de reprendre
directement les mesures du gabarit Figma « Slides ».

Usage CLI :
    python3 orem_pptx.py spec.json sortie.pptx
    python3 orem_pptx.py --demo sortie.pptx        # deck d'exemple, toutes les mises en page

Usage bibliothèque :
    from orem_pptx import build_deck
    build_deck(spec_dict, "sortie.pptx")

Format du spec (JSON) : {"slides": [ {"layout": "...", ...}, ... ]}
Mises en page : cover, chapter, intro, rows, bullets, cards (prix optionnel :
"price" + "price_unit"), steps, kpis,
quote, table, timeline, contact. Voir DEMO_SPEC plus bas pour chaque champ.

Polices : Trykker (titres) et Lexend Deca (corps). Elles doivent être
installées sur la machine qui ouvre le fichier (assets/fonts/ du skill),
sinon PowerPoint les remplace.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ASSETS = Path(__file__).resolve().parent.parent / "assets"
BG = ASSETS / "backgrounds"
LOGO = ASSETS / "logo" / "png"

PX = 6350  # EMU par pixel Figma (12 192 000 EMU / 1920 px)

# Couleurs
BLACK = "000000"
INK = "000012"
LIGHT = "F2F2F2"
WHITE = "FFFFFF"
PAPER = "F8F8F7"
ORANGE = "FF6200"
BLUE = "092482"
GREY_TXT = "434242"   # utilisé à 60 % d'opacité sur fond clair
SKY = "D2EDF6"

SERIF = "Trykker"
SANS = "Lexend Deca"
SANS_LIGHT = "Lexend Deca Light"
SANS_MEDIUM = "Lexend Deca Medium"
SANS_SEMIBOLD = "Lexend Deca SemiBold"


def px(v: float) -> Emu:
    return Emu(int(round(v * PX)))


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.upper())


def _set_alpha(run_or_fill_el, alpha: float):
    """Ajoute une transparence (0–1) à la couleur solide d'un élément rPr/solidFill."""
    srgb = run_or_fill_el.find(".//" + qn("a:srgbClr"))
    if srgb is None:
        return
    for old in srgb.findall(qn("a:alpha")):
        srgb.remove(old)
    a = srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
    srgb.append(a)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = px(1920)
        self.prs.slide_height = px(1080)
        self.blank = self.prs.slide_layouts[6]

    # --- primitives --------------------------------------------------------
    def slide(self, bg: str | None = BLACK, image: Path | None = None):
        s = self.prs.slides.add_slide(self.blank)
        if image is not None:
            s.shapes.add_picture(str(image), 0, 0, px(1920), px(1080))
        elif bg:
            fill = s.background.fill
            fill.solid()
            fill.fore_color.rgb = _rgb(bg)
        return s

    def text(self, s, x, y, w, h, content, *, font=SANS_LIGHT, size=24, color=LIGHT,
             alpha=None, tracking=-0.02, line=1.2, align="left", anchor="top",
             space_after=0, bullets=False):
        """Ajoute une zone de texte. `content` : str (\n = nouveau paragraphe) ou liste."""
        tb = s.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = None
        for side in ("left", "right", "top", "bottom"):
            setattr(tf, f"margin_{side}", 0)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        paras = content if isinstance(content, list) else str(content).split("\n")
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = line
            p.space_after = Pt(space_after * 0.5)
            if bullets:
                ptxt = "•  " + ptxt
            r = p.add_run()
            r.text = ptxt
            f = r.font
            f.name = font
            f.size = Pt(size * 0.5)
            f.color.rgb = _rgb(color)
            rpr = r._r.get_or_add_rPr()
            rpr.set("spc", str(int(round(size * 0.5 * tracking * 100))))
            # police est-asiatique/complexe identiques pour éviter les substitutions
            for tag in ("a:latin", "a:ea", "a:cs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = rpr.makeelement(qn(tag), {})
                    rpr.append(el)
                el.set("typeface", font)
            if alpha is not None:
                _set_alpha(rpr, alpha)
        return tb

    def rect(self, s, x, y, w, h, color, *, radius=None, alpha=None, line=None, shadow=False):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   px(x), px(y), px(w), px(h))
        if radius:
            shape.adjustments[0] = min(0.5, radius / min(w, h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(color)
        if alpha is not None:
            _set_alpha(shape.fill._xPr, alpha)
        if line:
            shape.line.color.rgb = _rgb(line[0])
            shape.line.width = px(line[1])
            if len(line) > 2:
                _set_alpha(shape.line._get_or_add_ln(), line[2])
        else:
            shape.line.fill.background()
        # Retire le style de thème (sinon ombre/contour hérités selon l'application)
        style = shape._element.find(qn("p:style"))
        if style is not None:
            shape._element.remove(style)
        sppr = shape._element.spPr
        for old in sppr.findall(qn("a:effectLst")):
            sppr.remove(old)
        eff = sppr.makeelement(qn("a:effectLst"), {})
        if shadow:  # ombre carte Orem : 0 16 90 #000012 @15 %
            sh = eff.makeelement(qn("a:outerShdw"), {"blurRad": str(px(90)), "dist": str(px(16)),
                                                     "dir": "5400000", "algn": "t", "rotWithShape": "0"})
            clr = sh.makeelement(qn("a:srgbClr"), {"val": INK})
            clr.append(clr.makeelement(qn("a:alpha"), {"val": "15000"}))
            sh.append(clr)
            eff.append(sh)
        sppr.append(eff)
        return shape

    def line(self, s, x1, y1, x2, y2, color=INK, width=1.5, alpha=None):
        ln = s.shapes.add_connector(1, px(x1), px(y1), px(x2), px(y2))
        ln.line.color.rgb = _rgb(color)
        ln.line.width = px(width)
        if alpha is not None:
            _set_alpha(ln.line._get_or_add_ln(), alpha)
        return ln

    def logo(self, s, which="symbole", color="noir", x=1749, y=64, w=107):
        path = LOGO / f"orem-{which}-{color}.png"
        return s.shapes.add_picture(str(path), px(x), px(y), width=px(w))

    def notes(self, s, text):
        if text:
            s.notes_slide.notes_text_frame.text = text

    def save(self, out):
        self.prs.save(out)


# --- mises en page ------------------------------------------------------------

def l_cover(d: Deck, c):
    s = d.slide(image=BG / "couverture-hero.jpg")
    d.logo(s, "symbole", "noir")
    d.text(s, 64, 600, 1792, 420, c["title"], font=SERIF, size=94, color=LIGHT, tracking=-0.03,
           line=1.15, anchor="bottom")
    if c.get("subtitle"):
        d.text(s, 64, 64, 1200, 60, c["subtitle"], size=24, color=LIGHT, alpha=0.8)
    return s


def l_chapter(d: Deck, c):
    style = c.get("style", "blob")
    if style == "light":
        s = d.slide(WHITE)
        d.text(s, 64, 330, 1200, 40, c.get("eyebrow", ""), size=24, color=ORANGE)
        d.text(s, 64, 380, 1300, 300, c["title"], font=SERIF, size=94, color=INK, tracking=-0.03, line=1.15)
        d.logo(s, "symbole", "noir", x=1350, y=330, w=420)
        return s
    img = BG / ("blob-sombre.jpg" if style == "blob" else "degrade-bleu.png")
    s = d.slide(image=img)
    top = 340 if style == "blob" else 300
    d.text(s, 160, top, 1600, 40, c.get("eyebrow", ""), size=24, color=ORANGE, align="center")
    d.text(s, 160, top + 50, 1600, 300, c["title"], font=SERIF, size=94, color=LIGHT, tracking=-0.03,
           line=1.15, align="center")
    return s


def l_intro(d: Deck, c):
    dark = c.get("theme", "light") == "dark"
    s = d.slide(BLACK if dark else WHITE)
    fg = LIGHT if dark else INK
    sec = LIGHT if dark else GREY_TXT
    d.text(s, 64, 230, 860, 40, c.get("eyebrow", ""), size=24, color=ORANGE)
    d.text(s, 64, 280, 860, 330, c["title"], font=SERIF, size=72, color=fg, tracking=-0.03, anchor="bottom")
    if c.get("body"):
        d.text(s, 64, 660, 860, 150, c["body"], size=24, color=sec, alpha=0.6)
    if c.get("emphasis"):
        d.text(s, 64, 830, 860, 60, c["emphasis"], font=SANS_SEMIBOLD, size=24, color=sec, alpha=0.6)
    if c.get("image"):
        s.shapes.add_picture(str(c["image"]), px(972), 0, px(948), px(1080))
    else:
        s.shapes.add_picture(str(BG / "degrade-bleu.png"), px(972), 0, px(948), px(1080))
    return s


def l_rows(d: Deck, c):
    dark = c.get("theme", "dark") == "dark"
    s = d.slide(image=BG / ("trame-points-sombre.png" if dark else "trame-points-clair.png"))
    fg = LIGHT if dark else INK
    d.text(s, 64, 50, 1792, 120, c["title"], font=SERIF, size=94, color=fg, tracking=-0.03, line=1.1)
    y = 170
    if c.get("subtitle"):
        d.text(s, 64, 170, 1792, 60, c["subtitle"], font=SERIF, size=37, color=fg, alpha=0.6, tracking=0)
        y = 250
    rows = c["rows"]
    step = min(150, (1016 - y - 40) / max(1, len(rows)))
    y += 40
    for label, value in rows:
        d.text(s, 64, y, 560, step, label, font=SERIF, size=37, color=fg, alpha=0.6, tracking=0, line=1.3)
        d.text(s, 684, y, 1172, step, value, font=SERIF, size=37, color=fg, tracking=0, line=1.3)
        y += step
    return s


def l_bullets(d: Deck, c):
    dark = c.get("theme", "dark") == "dark"
    s = d.slide(image=BG / ("trame-points-sombre.png" if dark else "trame-points-clair.png"))
    fg = LIGHT if dark else INK
    d.text(s, 64, 50, 1792, 120, c["title"], font=SERIF, size=94, color=fg, tracking=-0.03, line=1.1)
    if c.get("subtitle"):
        d.text(s, 64, 170, 1792, 60, c["subtitle"], font=SERIF, size=37, color=fg, alpha=0.6, tracking=0)
    cols = c["columns"]
    w = (1792 - 64 * (len(cols) - 1)) / len(cols)
    for i, col in enumerate(cols):
        x = 64 + i * (w + 64)
        d.text(s, x, 330, w, 60, col.get("label", ""), font=SERIF, size=37, color=fg, alpha=0.6, tracking=0)
        d.text(s, x, 420, w, 600, col["items"], font=SERIF, size=37, color=fg, tracking=0, line=1.15,
               bullets=True, space_after=6)
    return s


def l_cards(d: Deck, c):
    s = d.slide(image=BG / "degrade-bleu.png")
    cards = c["cards"]
    n = len(cards)
    gap = 64
    w = (1792 - gap * (n - 1)) / n
    top = 250
    if c.get("title"):
        d.text(s, 64, 80, 1792, 120, c["title"], font=SERIF, size=64, color=LIGHT, tracking=-0.03)
    for i, card in enumerate(cards):
        x = 64 + i * (w + gap)
        d.rect(s, x, top, w, 520, WHITE, radius=24, shadow=True)
        d.text(s, x + 40, top + 40, w - 80, 30, card.get("tag", f"Opt {i + 1}"), font=SANS, size=20, color=ORANGE)
        d.text(s, x + 40, top + 80, w - 80, 110, card["title"], font=SERIF, size=37, color=INK, tracking=0, line=1.3)
        y_items = top + 190
        if card.get("price"):  # ex. "290 €" + price_unit "HT / mois"
            d.text(s, x + 40, top + 150, w - 80, 90, card["price"], font=SERIF, size=64, color=INK, tracking=-0.03,
                   line=1.1)
            if card.get("price_unit"):
                d.text(s, x + 40, top + 240, w - 80, 30, card["price_unit"], font=SANS_LIGHT, size=20, color=GREY_TXT,
                       alpha=0.6)
            y_items = top + 290
        if card.get("items"):
            d.text(s, x + 40, y_items, w - 80, top + 400 - y_items, card["items"], font=SANS, size=24, color=INK,
                   bullets=True, space_after=4)
        if card.get("note"):
            d.text(s, x + 40, top + 400, w - 80, 80, card["note"], font=SANS, size=20, color=GREY_TXT, alpha=0.6,
                   anchor="bottom")
    return s


def l_steps(d: Deck, c):
    s = d.slide(WHITE)
    s.shapes.add_picture(str(BG / "degrade-bleu-vertical.png"), 0, 0, px(653), px(1080))
    d.text(s, 24, 300, 605, 480, c["title"], font=SANS, size=64, color=LIGHT, tracking=-0.03, anchor="middle")
    steps = c["steps"][:4]
    for i, st in enumerate(steps):
        x = 869 + (i % 2) * 452
        y = 330 + (i // 2) * 250
        d.text(s, x, y, 420, 30, st.get("tag", f"Étape {i + 1:02d}"), font=SANS, size=17, color=ORANGE)
        d.text(s, x, y + 45, 420, 60, st["title"], font=SERIF, size=37, color=INK, tracking=0, line=1.3)
        d.text(s, x, y + 110, 420, 110, st.get("text", ""), font=SANS, size=24, color=INK)
    return s


def l_kpis(d: Deck, c):
    s = d.slide(WHITE)
    if c.get("caption"):
        d.text(s, 64, 64, 700, 40, c["caption"], font=SANS, size=20, color=INK)
    y = 64
    for k in c["kpis"][:4]:
        d.line(s, 928, y, 928, y + 105, INK, 1.5)
        d.text(s, 992, y - 5, 420, 115, k["value"], font=SERIF, size=94, color=INK, tracking=-0.03, line=1.1)
        d.text(s, 1400, y + 15, 460, 100, k["label"], font=SERIF, size=37, color=INK, tracking=0, line=1.2)
        y += 170
    q = c.get("quote")
    if q:
        d.text(s, 64, 700, 620, 220, f"“{q['text']}”", font=SANS, size=37, color=INK, line=1.3, tracking=0,
               anchor="bottom")
        d.text(s, 64, 960, 620, 40, q.get("author", ""), font=SANS, size=20, color=INK)
    return s


def l_quote(d: Deck, c):
    style = c.get("style", "gradient")
    s = d.slide(image=BG / ("degrade-bleu.png" if style == "gradient" else "blob-sombre.jpg"))
    d.text(s, 400, 300, 1120, 400, f"“{c['text']}”", font=SERIF, size=72, color=WHITE, tracking=-0.03,
           anchor="bottom")
    if c.get("author"):
        d.text(s, 400, 720, 1120, 60, c["author"], font=SERIF, size=37, color=WHITE, align="right", tracking=0)
    return s


def l_table(d: Deck, c):
    s = d.slide(WHITE)
    d.text(s, 64, 64, 560, 200, c["title"], font=SERIF, size=64, color=INK, tracking=-0.03)
    cols = c["columns"]
    rows = c["rows"]
    x0, w = 684, 1172
    cw = w / len(cols)
    for j, h in enumerate(cols):
        d.text(s, x0 + 16 + j * cw, 64, cw - 32, 40, h, font=SANS_LIGHT, size=20, color=GREY_TXT, alpha=0.6)
    d.line(s, x0, 124, x0 + w, 124, "BCBAB9", 1, alpha=0.5)
    rh = min(52, (900 - 160) / max(1, len(rows)))
    y = 160
    for i, r in enumerate(rows):
        if i % 2 == 0:
            d.rect(s, x0, y, w, rh, PAPER)
        for j, cell in enumerate(r):
            d.text(s, x0 + 16 + j * cw, y, cw - 32, rh, str(cell), font=SANS, size=20, color=INK, anchor="middle")
        y += rh
    if c.get("note"):
        d.text(s, 64, 900, 520, 120, c["note"], font=SANS, size=20, color=GREY_TXT, alpha=0.6, anchor="bottom")
    return s


def l_timeline(d: Deck, c):
    s = d.slide(WHITE)
    d.text(s, 296, 64, 1400, 200, c["title"], font=SERIF, size=64, color=INK, tracking=-0.03)
    steps = c["steps"][:4]
    d.line(s, 320, 506, 1920, 506, INK, 1.5)
    for i, st in enumerate(steps):
        x = 310 + i * 406
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, px(x), px(496), px(20), px(20))
        dot.fill.solid()
        dot.fill.fore_color.rgb = _rgb(INK)
        dot.line.fill.background()
        d.text(s, x, 560, 360, 110, st["title"], font=SERIF, size=37, color=INK, tracking=0, line=1.3)
        d.text(s, x, 690, 300, 120, st.get("text", ""), font=SANS_LIGHT, size=20, color=GREY_TXT, alpha=0.6)
    return s


def l_contact(d: Deck, c):
    s = d.slide(image=BG / "contact-halo.jpg")
    d.rect(s, 705, 216, 474, 598, WHITE, radius=24, alpha=0.05, line=(LIGHT, 1, 0.2))
    d.text(s, 745, 256, 400, 30, c.get("eyebrow", "Contact"), font=SANS, size=17, color=ORANGE)
    d.text(s, 745, 300, 400, 60, c["name"], font=SERIF, size=37, color=LIGHT, tracking=0)
    d.text(s, 745, 480, 400, 120, c.get("lines", []), size=20, color=LIGHT, alpha=0.6, space_after=10)
    if c.get("cta"):
        btn = d.rect(s, 745, 700, 240, 54, WHITE, radius=27)
        tf = btn.text_frame
        tf.text = c["cta"]  # Lexend Deca n'a pas de glyphe flèche : pas de ↗ dans les .pptx
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.runs[0]
        r.font.name, r.font.size, r.font.color.rgb = SANS, Pt(10), _rgb(INK)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return s


LAYOUTS = {
    "cover": l_cover, "chapter": l_chapter, "intro": l_intro, "rows": l_rows, "bullets": l_bullets,
    "cards": l_cards, "steps": l_steps, "kpis": l_kpis, "quote": l_quote, "table": l_table,
    "timeline": l_timeline, "contact": l_contact,
}


def build_deck(spec: dict, out: str | Path):
    d = Deck()
    for i, c in enumerate(spec["slides"]):
        fn = LAYOUTS.get(c.get("layout"))
        if fn is None:
            raise ValueError(f"slide {i + 1} : mise en page inconnue {c.get('layout')!r} ({', '.join(LAYOUTS)})")
        s = fn(d, c)
        d.notes(s, c.get("notes"))
    d.save(str(out))
    return out


DEMO_SPEC = {
    "slides": [
        {"layout": "cover", "title": "Recommandation stratégique & opérationnelle pour le futur site Client"},
        {"layout": "chapter", "style": "blob", "eyebrow": "Programme partenaire",
         "title": "Transformer son réseau\nen valeur concrète"},
        {"layout": "intro", "eyebrow": "Notre approche", "title": "Des sites conçus pour faire travailler les entreprises B2B.",
         "body": "Nos clients ont une offre solide et un site en retard sur leur niveau. On commence par le fond avant de concevoir quoi que ce soit.",
         "emphasis": "Recommander Orem, c’est recommander quelqu’un qui livre."},
        {"layout": "rows", "title": "Ce que nous avons compris du projet",
         "subtitle": "Le client entre dans une nouvelle phase, où le site doit permettre de :",
         "rows": [["Positionner", "Aligner l’image digitale avec le niveau réel d’expertise."],
                  ["Convaincre", "Créer un outil utile en acquisition, en réassurance et en rendez-vous commercial."],
                  ["Déployer", "Construire une première phase forte, puis étendre progressivement le site."]]},
        {"layout": "bullets", "title": "Notre lecture éditoriale", "columns": [
            {"label": "Faire comprendre avant d’impressionner", "items": ["Expertise", "Clarté", "Sobriété", "Crédibilité"]},
            {"label": "Notre objectif", "items": ["Une parole plus directe", "Davantage de bénéfices concrets", "Plus de preuves"]}]},
        {"layout": "chapter", "style": "gradient", "eyebrow": "Programme partenaire", "title": "Une récompense au choix"},
        {"layout": "cards", "cards": [
            {"tag": "Opt 1", "title": "Commission cash", "price": "12 %", "price_unit": "puis 17 % dès le 3e deal", "items": ["12 % sur les 2 premiers deals", "17 % à partir du 3e"], "note": "Versé dès encaissement du premier paiement client"},
            {"tag": "Opt 2", "title": "Production offerte", "items": ["2 mois d’accompagnement offerts"], "note": "Utilisables pour faire évoluer le site de ses clients"},
            {"tag": "Opt 3", "title": "Avantage sur son site", "items": ["25 % de réduction sur la refonte"], "note": "Un retour direct sur son image"}]},
        {"layout": "steps", "title": "Une introduction.\nOrem gère le reste.", "steps": [
            {"title": "Identification", "text": "Une entreprise du réseau correspond au profil cible."},
            {"title": "Mise en relation", "text": "Un email d’introduction suffit."},
            {"title": "Orem gère", "text": "Qualification, devis, signature, encaissement."},
            {"title": "Récompense", "text": "Deal validé, la récompense est activée."}]},
        {"layout": "kpis", "caption": "Les chiffres parlent d’eux-mêmes.", "kpis": [
            {"value": "20 000", "label": "€ de CA générés dès la première semaine"},
            {"value": "×2", "label": "du MRR de la plateforme en 48 h"},
            {"value": "5 119", "label": "sessions en un seul mois"}],
         "quote": {"text": "Le boss pour les sites et la meilleure personne pour vous aider.", "author": "Prénom Nom — Entreprise"}},
        {"layout": "quote", "text": "On préfère refuser un projet que livrer quelque chose qui ne correspond pas.", "author": "Orem"},
        {"layout": "table", "title": "Accompagnement post-livraison", "columns": ["Chantier", "Client", "Orem"],
         "rows": [["Architecture", "Validation", "Conception"], ["Direction artistique", "Retours", "Production"],
                  ["UX design", "Ateliers", "Animation"], ["SEO / GEO", "Mots-clés", "Stratégie"]],
         "note": "Le chemin critique dépend autant de la production que des validations client."},
        {"layout": "timeline", "title": "Une introduction,\nOrem gère le reste.", "steps": [
            {"title": "Identification", "text": "Une entreprise du réseau correspond au profil cible."},
            {"title": "Mise en relation", "text": "Un email d’introduction suffit."},
            {"title": "Récompense", "text": "Deal validé, la récompense est activée."},
            {"title": "Orem gère", "text": "Qualification, devis, signature."}]},
        {"layout": "contact", "name": "Prénom Nom", "lines": ["prenom@agence-orem.fr", "+33 6 00 00 00 00"], "cta": "Prochaine étape"},
    ]
}


def main(argv):
    if len(argv) == 3 and argv[1] == "--demo":
        build_deck(DEMO_SPEC, argv[2])
    elif len(argv) == 3:
        build_deck(json.loads(Path(argv[1]).read_text(encoding="utf-8")), argv[2])
    else:
        print(__doc__)
        return 1
    print(f"OK → {argv[2]}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
